"""
O'QUV YUKLAMASI HAJMI ko'rinishidagi Excel fayllarni o'qish algoritmi.

Fayl tuzilishi (ustunlar 1-indeksli):
    A(1)  - T/r (tartib raqami)
    B-D   - Fan nomi (birlashtirilgan / merged)
    E(5)  - Ta'lim tili
    F(6)  - Semestr
    G(7)  - Talabalar soni
    H(8)  - Guruhlar soni
    I(9)  - Patoklar
    J(10) - Ma'ruza: Reja
    K(11) - Ma'ruza: Jami
    L(12) - Amaliy mashg'ulot: Reja
    M(13) - Amaliy mashg'ulot: Jami
    N(14) - Labratoriya mashg'uloti: Reja
    O(15) - Labratoriya mashg'uloti: Jami
    P(16) - Nazorat
    Q(17) - Kurs loyihasi
    R(18) - MD ilmiy rahbarlik
    S(19) - BMI rahbarlik
    T(20) - Amaliyot
    U(21) - DAK, Taqriz
    V(22) - Jami

Jadval ichida bo'lim ajratuvchi qatorlar bo'ladi:
    "Kuzgi semestr"                 -> kuzgi bo'lim boshlanadi
    "Kuzgi semestr bo'yicha jami:"  -> kuzgi yakuniy (jami) qatori
    "Bahorgi semestr"               -> bahorgi bo'lim boshlanadi
    "Bahorgi semestr bo'yicha jami:"-> bahorgi yakuniy (jami) qatori
    "O'quv yili bo'yicha jami:"     -> o'quv yili yakuniy qatori
Undan keyingi PO'T / O'YX (xodimlar) bo'limi o'qilmaydi.
"""

import logging
import math
from pathlib import Path

from openpyxl import load_workbook
import xlrd

logger = logging.getLogger(__name__)


# Ustun indekslari (1-indeksli) -> chiqish kalitlari
COLUMN_MAP = {
    1: 'tr',
    5: 'talim_tili',
    6: 'semestr',
    7: 'talabalar_soni',
    8: 'guruhlar_soni',
    9: 'patok',
    10: 'maruza_reja',
    11: 'maruza_jami',
    12: 'amaliy_reja',
    13: 'amaliy_jami',
    14: 'laboratoriya_reja',
    15: 'laboratoriya_jami',
    16: 'nazorat',
    17: 'kurs_loyihasi',
    18: 'md_ilmiy_rahbarlik',
    19: 'bmi_rahbarlik',
    20: 'amaliyot',
    21: 'dak_taqriz',
    22: 'jami',
}

# Fan nomi birlashtirilgan ustunlar (B, C, D)
FAN_NOMI_COLS = (2, 3, 4)

# Matn ustunlari (raqamga aylantirilmaydi)
TEXT_FIELDS = {'talim_tili'}


def _round_custom(x):
    """Maxsus yaxlitlash: .0–.5 -> pastga, .6–.9 -> yuqoriga.
    Masalan 5.1–5.5 -> 5, 5.6–5.9 -> 6."""
    return int(math.floor(x + 0.4))


def _to_number(value):
    """Katak qiymatini butun songa (yaxlitlab) keltiradi. Matn bo'lsa o'zicha qaytaradi."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return _round_custom(float(value))
    text = str(value).strip()
    if text == '':
        return None
    # Vergulli o'nlik sonlarni ham qo'llab-quvvatlaymiz (275,00)
    normalized = text.replace(' ', '').replace(',', '.')
    try:
        return _round_custom(float(normalized))
    except ValueError:
        return text


def _cell_text(value):
    """Katak qiymatini tozalangan matnga aylantiradi."""
    if value is None:
        return ''
    return str(value).strip()


def _row_text(cells):
    """Qatordagi barcha matnli kataklarni bitta kichik harfli satrga birlashtiradi."""
    parts = [_cell_text(c).lower() for c in cells if _cell_text(c)]
    return ' '.join(parts)


def _fan_nomi(cells):
    """B, C, D ustunlaridan fan nomini yig'adi. Birlashtirilgan (merged) kataklar
    bir xil qiymatni takrorlashi mumkin — takroriy bo'laklarni olib tashlaymiz."""
    parts = []
    for col in FAN_NOMI_COLS:
        text = _cell_text(cells[col - 1]) if len(cells) >= col else ''
        if text and text not in parts:
            parts.append(text)
    # Ko'p bo'sh joylarni bittaga keltiramiz (dedup uchun)
    return ' '.join(' '.join(parts).split())


def _is_total_row(row_text):
    return 'jami' in row_text


def _detect_section(row_text):
    """Bo'lim ajratuvchi qatorni aniqlaydi. None bo'lsa — bu marker emas."""
    if 'semestr' not in row_text:
        return None
    if 'kuzgi' in row_text:
        return 'kuzgi'
    if 'bahorgi' in row_text or 'bahor' in row_text:
        return 'bahorgi'
    return None


def _detect_education_type(row_text):
    """Ta'lim shakli ajratuvchi qatorni frontend qiymatiga aylantiradi."""
    if 'ta’lim' not in row_text and "ta'lim" not in row_text and 'talim' not in row_text:
        return None
    if 'kunduzgi' in row_text:
        return 'bakalavr'
    # "Sirtqi va 2-OT ta'lim shakli" sirtqi sifatida ko'rsatiladi.
    if 'sirtqi' in row_text:
        return 'sirtqi'
    if 'masofaviy' in row_text or 'masofa' in row_text:
        return 'masofaviy'
    if 'magistratura' in row_text or 'magistr' in row_text:
        return 'magistr'
    if "qo'shma" in row_text or 'qo‘shma' in row_text or 'qoshma' in row_text:
        return '2-ot'
    return None


def _detect_result_columns(rows):
    """Sarlavhadan DAK/Taqriz va Jami ustunlarining 0-indeksini aniqlaydi."""
    dak_index = None
    jami_index = None
    for cells in rows:
        row_text = _row_text(cells)
        if _detect_section(row_text):
            break
        for index, value in enumerate(cells):
            text = _cell_text(value).lower().strip().rstrip(':')
            if text == 'jami':
                jami_index = index
            if 'dak' in text or 'taqriz' in text:
                dak_index = index

    # Sarlavha topilmagan fayllar uchun eski tuzilishga mos fallback.
    column_count = max((len(row) for row in rows), default=0)
    if jami_index is None:
        jami_index = 21 if column_count >= 22 else 20 if column_count >= 21 else None
    if dak_index is None and jami_index == 21:
        dak_index = 20
    return dak_index, jami_index


def _parse_row(cells, dak_index=None, jami_index=None):
    """Bitta data qatorini lug'atga aylantiradi."""
    item = {}
    for col, key in COLUMN_MAP.items():
        if key == 'dak_taqriz':
            raw = cells[dak_index] if dak_index is not None and len(cells) > dak_index else None
        elif key == 'jami':
            raw = cells[jami_index] if jami_index is not None and len(cells) > jami_index else None
        else:
            raw = cells[col - 1] if len(cells) >= col else None
        if key in TEXT_FIELDS:
            item[key] = _cell_text(raw) or None
        else:
            item[key] = _to_number(raw)
    item['fan_nomi'] = _fan_nomi(cells)
    return item


def _is_data_row(item, row_text):
    """Data qatori bo'lsa True. Bo'sh yoki sarlavha qatorlarini chiqarib tashlaydi."""
    if not item['fan_nomi']:
        return False
    # Kamida bitta raqamli ko'rsatkich yoki tartib raqami bo'lishi kerak
    has_number = item['tr'] is not None or any(
        item[k] is not None
        for k in COLUMN_MAP.values()
        if k not in ('tr',) and k not in TEXT_FIELDS
    )
    return has_number


def _calculate_control(student_count, education_type, subject_name=''):
    """Nazorat = talabalar soni × ta'lim shakli koeffitsienti."""
    normalized_name = _cell_text(subject_name).lower()
    if (
        'ilmiy pedagogik ish' in normalized_name
        or 'magistr dissertatsiyasiga rahbarlik' in normalized_name
    ):
        return 0
    if isinstance(student_count, bool) or not isinstance(student_count, (int, float)):
        return 0
    coefficient = 1 if education_type == 'sirtqi' else 0.5
    value = float(student_count) * coefficient
    # Talabalar soni musbat bo'ladi: .5 va undan yuqori natijani tepaga
    # yaxlitlaymiz (masalan, 45 × 0.5 = 22.5 -> 23).
    return int(math.floor(value + 0.5))


def _read_modern_excel_rows(file_obj):
    """`.xlsx` va `.xlsm` fayllarining birinchi varag'ini o'qiydi."""
    wb = load_workbook(file_obj, data_only=True)
    try:
        ws = wb.active

        merged_lookup = {}
        for rng in ws.merged_cells.ranges:
            anchor_value = ws.cell(row=rng.min_row, column=rng.min_col).value
            if anchor_value is None:
                continue
            for row_index in range(rng.min_row, rng.max_row + 1):
                for column_index in range(rng.min_col, rng.max_col + 1):
                    merged_lookup[(row_index, column_index)] = anchor_value

        rows = []
        for row_index, values in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = list(values or ())
            for column_index in range(len(cells)):
                merged_value = merged_lookup.get((row_index, column_index + 1))
                if merged_value is not None:
                    cells[column_index] = merged_value
            rows.append(cells)
        return rows
    finally:
        wb.close()


def _read_legacy_excel_rows(file_obj):
    """Eski binary `.xls` faylining birinchi varag'ini o'qiydi."""
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    content = file_obj.read()
    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        sheet = book.sheet_by_index(0)
        merged_lookup = {}
        for row_start, row_end, col_start, col_end in sheet.merged_cells:
            anchor_value = sheet.cell_value(row_start, col_start)
            if anchor_value in (None, ''):
                continue
            for row_index in range(row_start, row_end):
                for column_index in range(col_start, col_end):
                    merged_lookup[(row_index, column_index)] = anchor_value

        rows = []
        for row_index in range(sheet.nrows):
            cells = sheet.row_values(row_index, start_colx=0, end_colx=sheet.ncols)
            for column_index in range(len(cells)):
                merged_value = merged_lookup.get((row_index, column_index))
                if merged_value is not None:
                    cells[column_index] = merged_value
            rows.append(cells)
        return rows
    finally:
        book.release_resources()


def _read_excel_rows(file_obj):
    """Fayl kengaytmasiga qarab mos Excel o'quvchisini tanlaydi."""
    extension = Path(getattr(file_obj, 'name', '')).suffix.lower()
    if extension == '.xls':
        return _read_legacy_excel_rows(file_obj)
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    return _read_modern_excel_rows(file_obj)


def parse_workload_excel(file_obj):
    """
    O'quv yuklamasi Excel faylini o'qib, semestrlarga ajratilgan ma'lumot qaytaradi.

    file_obj — fayl yo'li yoki fayl-ga o'xshash obyekt (request.FILES[...]).

    Qaytaradi:
        {
            'meta': {'title': '...'},
            'kuzgi': [ {...}, ... ],
            'bahorgi': [ {...}, ... ],
            'totals': {
                'kuzgi': {...} | None,
                'bahorgi': {...} | None,
                'yillik': {...} | None,
            },
            'warnings': [...],
        }
    """
    rows = _read_excel_rows(file_obj)
    dak_index, jami_index = _detect_result_columns(rows)

    result = {
        'meta': {'title': ''},
        'kuzgi': [],
        'bahorgi': [],
        'totals': {'kuzgi': None, 'bahorgi': None, 'yillik': None},
        'warnings': [],
    }

    section = None  # 'kuzgi' | 'bahorgi' | None
    education_type = 'bakalavr'
    title_parts = []
    seen_first_marker = False

    for cells in rows:
        if cells is None:
            continue
        cells = list(cells)
        text = _row_text(cells)
        if not text:
            continue

        # Jadval ichidagi ta'lim shakli sarlavhasini keyingi fan qatorlari
        # uchun eslab qolamiz. Sarlavhaning o'zi data qatori emas.
        detected_education_type = _detect_education_type(text)
        if detected_education_type:
            education_type = detected_education_type
            continue

        # Sarlavha (yuqoridagi) qatorlardan title yig'amiz
        if not seen_first_marker and ('yuklamasi' in text or 'kafedra' in text or "o‘quv" in text or 'oquv' in text):
            joined = ' '.join(_cell_text(c) for c in cells if _cell_text(c))
            if joined:
                title_parts.append(joined)

        # 1) Yakuniy (jami) qatorlarini ushlaymiz
        if _is_total_row(text):
            totals = _parse_row(cells, dak_index, jami_index)
            if 'yili' in text:
                result['totals']['yillik'] = totals
                section = None  # o'quv yili jamisidan keyin data tugaydi
            elif 'kuzgi' in text:
                result['totals']['kuzgi'] = totals
            elif 'bahorgi' in text or 'bahor' in text:
                result['totals']['bahorgi'] = totals
            continue

        # 2) Bo'lim ajratuvchi marker
        detected = _detect_section(text)
        if detected:
            section = detected
            seen_first_marker = True
            continue

        # 3) Data qatori (faqat aktiv semestr bo'limi ichida)
        if section is None:
            continue

        item = _parse_row(cells, dak_index, jami_index)
        if _is_data_row(item, text):
            item['talim_shakli'] = education_type
            item['nazorat'] = _calculate_control(
                item.get('talabalar_soni'), education_type, item.get('fan_nomi'),
            )
            result[section].append(item)
            logger.info("[Excel] %s qatori o'qildi: %s", section, item)

    result['meta']['title'] = ' | '.join(title_parts[:2])

    logger.info(
        "[Excel] Fayl o'qildi. title=%r | kuzgi=%d ta | bahorgi=%d ta | totals=%s | warnings=%s",
        result['meta']['title'],
        len(result['kuzgi']),
        len(result['bahorgi']),
        result['totals'],
        result['warnings'],
    )

    if not result['kuzgi'] and not result['bahorgi']:
        result['warnings'].append(
            "Hech qanday data qatori topilmadi. Fayl tuzilishi kutilgan "
            "O'QUV YUKLAMASI HAJMI ko'rinishida ekanligini tekshiring."
        )

    return result
